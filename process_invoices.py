#!/usr/bin/env python3
import sys
import subprocess
import importlib
from pathlib import Path

REQUIRED_PACKAGES = [
    ("pandas", ["pandas"]),
    ("openpyxl", ["openpyxl"]),
    ("PIL", ["pillow"]),
    ("numpy", ["numpy"]),
    ("dateutil", ["python-dateutil"]),
    ("regex", ["regex"]),
]


TORCH_INSTALL_COMMAND = [
    sys.executable,
    "-m",
    "pip",
    "install",
    "torch",
    "torchvision",
    "torchaudio",
    "--index-url",
    "https://download.pytorch.org/whl/cpu",
]


def ensure_packages():
    for module_name, packages in REQUIRED_PACKAGES:
        try:
            importlib.import_module(module_name)
        except ModuleNotFoundError:
            subprocess.check_call([sys.executable, "-m", "pip", "install", *packages])
    try:
        importlib.import_module("torch")
    except ModuleNotFoundError:
        subprocess.check_call(TORCH_INSTALL_COMMAND)
    try:
        importlib.import_module("easyocr")
    except ModuleNotFoundError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "easyocr"])


ensure_packages()

import csv
import statistics
import regex
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

import numpy as np  # noqa: F401
import pandas as pd  # noqa: F401
import dateutil  # noqa: F401
from PIL import Image  # noqa: F401
import easyocr
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# Settings
DRY_RUN = False
TAX_RATE_DEFAULT = 3.3125
TEMPLATE_NAMES = [
    "Breakdown charges invoice summary",
    "Breakdown charges- Invoices Summary",
]

ROOT = Path(__file__).resolve().parent
PO_IMAGE_PATH = ROOT / "po_screen.png"
INVOICE_IMAGE_PATH = ROOT / "pvault_invoice_log.png"
TEMPLATE_PATH = ROOT / "Breakdown charges- Invoices Summary.xlsx"
OUTPUT_PATH = ROOT / "Breakdown charges- Invoices Summary POPULATED.xlsx"
CSV_PATH = ROOT / "ocr_invoice_rows.csv"
LOG_PATH = ROOT / "build_log.txt"

CURRENCY_FORMAT = "$#,##0.00"
TEXT_FORMAT = "@"
MATRIX_FILL = PatternFill(start_color="FFF5F5F5", end_color="FFF5F5F5", fill_type="solid")
VARIANCE_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")


@dataclass
class OCRItem:
    text: str
    confidence: float
    x_center: float
    y_center: float
    x_min: float
    x_max: float


@dataclass
class InvoiceRow:
    invoice_number: str
    po_line: Optional[int]
    gross: Optional[float]
    tax: Optional[float]
    units: str
    unit_cost: str
    status: str
    description: str
    confidence: float


class WorkbookContext:
    def __init__(self):
        self.template_used = False
        self.template_created = False
        self.workbook = None
        self.sheet = None


class BuildLogger:
    def __init__(self):
        self.messages: List[str] = []

    def log(self, message: str):
        print(f"[LOG] {message}")
        self.messages.append(message)

    def write(self, path: Path):
        path.write_text("\n".join(self.messages), encoding="utf-8")


def normalize_text(text: str) -> str:
    if text is None:
        return ""
    return " ".join(text.replace("\xa0", " ").split())


def normalize_for_match(text: str) -> str:
    return regex.sub(r"\s+", "", normalize_text(text).casefold())


def load_or_create_workbook(logger: BuildLogger) -> WorkbookContext:
    ctx = WorkbookContext()
    if TEMPLATE_PATH.exists():
        logger.log("Template workbook found. Loading existing file.")
        wb = load_workbook(TEMPLATE_PATH)
        target_sheet = None
        for name in TEMPLATE_NAMES:
            if name in wb.sheetnames:
                target_sheet = wb[name]
                logger.log(f"Using sheet '{name}' as template sheet.")
                break
        if target_sheet is None:
            for sheet in wb.worksheets:
                if sheet.sheet_state == "visible":
                    target_sheet = sheet
                    logger.log(
                        f"No sheet matched template names. Using first visible sheet '{sheet.title}'."
                    )
                    break
        ctx.workbook = wb
        ctx.sheet = target_sheet if target_sheet is not None else wb.active
        ctx.template_used = True
        wb.save(OUTPUT_PATH)
        logger.log("Saved copy as populated workbook for editing.")
    else:
        logger.log("Template workbook not found. Creating new workbook with required layout.")
        wb = Workbook()
        ws = wb.active
        ws.title = "Breakdown charges- Invoices Summary"
        ws.freeze_panes = "A24"
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 16
        ws.cell(row=1, column=4, value="Tax Rate")
        headers = [
            "Cable Size",
            "bom",
            "Qty Ordered",
            "Unit Cost",
            "Qty Delivered (Vendor)",
            "Delta PRE/WSE",
            "Delta to Ordered",
            "Total",
            "Credit/Debit",
            "Notes",
        ]
        for idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=idx, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in range(4, 22):
            for col in range(1, 11):
                ws.cell(row=row, column=col)
        ws.cell(row=24, column=1, value="INV")
        for row in range(25, 43):
            for col in range(1, 11):
                ws.cell(row=row, column=col)
        ws.cell(row=44, column=1, value="Total by invoice")
        ws.cell(row=45, column=1, value="Invoice total from log")
        ws.cell(row=46, column=1, value="Variance")
        ctx.workbook = wb
        ctx.sheet = ws
        ctx.template_created = True
        wb.save(OUTPUT_PATH)
        logger.log("Created new workbook and saved populated copy for editing.")
    return ctx


def read_image_text(image_path: Path, reader: easyocr.Reader, logger: BuildLogger):
    if not image_path.exists():
        raise FileNotFoundError(f"Required image not found: {image_path}")
    logger.log(f"Running OCR on {image_path.name}.")
    return reader.readtext(str(image_path), detail=1, paragraph=False)


def convert_results_to_items(results) -> List[OCRItem]:
    items: List[OCRItem] = []
    for bbox, text, confidence in results:
        cleaned = normalize_text(text)
        x_values = [point[0] for point in bbox]
        y_values = [point[1] for point in bbox]
        items.append(
            OCRItem(
                text=cleaned,
                confidence=float(confidence),
                x_center=float(sum(x_values) / len(x_values)),
                y_center=float(sum(y_values) / len(y_values)),
                x_min=float(min(x_values)),
                x_max=float(max(x_values)),
            )
        )
    return items


def group_items_by_row(items: List[OCRItem], tolerance: float = 12.0) -> List[List[OCRItem]]:
    rows: List[List[OCRItem]] = []
    centers: List[float] = []
    for item in sorted(items, key=lambda it: it.y_center):
        assigned = False
        for idx, center in enumerate(centers):
            if abs(item.y_center - center) <= tolerance:
                rows[idx].append(item)
                centers[idx] = sum(x.y_center for x in rows[idx]) / len(rows[idx])
                assigned = True
                break
        if not assigned:
            rows.append([item])
            centers.append(item.y_center)
    for row in rows:
        row.sort(key=lambda it: it.x_center)
    return rows


def find_po_header_index(rows: List[List[OCRItem]]) -> Optional[int]:
    for idx, row in enumerate(rows):
        text = " ".join(item.text.casefold() for item in row)
        if ("po item" in text or "po line" in text) and "description" in text:
            return idx
    return None


def compute_description_bounds(header_row: List[OCRItem], fallback_max: float) -> Tuple[float, float]:
    desc_tokens = [item for item in header_row if "description" in item.text.casefold()]
    item_tokens = [item for item in header_row if "item" in item.text.casefold()]
    units_tokens = [item for item in header_row if "unit" in item.text.casefold()]
    if desc_tokens:
        left = min(item.x_min for item in desc_tokens)
        right = max(item.x_max for item in desc_tokens)
    elif item_tokens:
        left = min(item.x_min for item in item_tokens)
        right = max(item.x_max for item in item_tokens)
    else:
        left = min(item.x_min for item in header_row)
        right = fallback_max
    if units_tokens:
        units_left = min(item.x_min for item in units_tokens)
        right = (right + units_left) / 2
    return left, max(right, left + 1.0)


def extract_po_lines(rows: List[List[OCRItem]], logger: BuildLogger) -> Tuple[List[int], Dict[int, str]]:
    header_idx = find_po_header_index(rows)
    if header_idx is None:
        logger.log("PO header row not found; treating all rows as candidates.")
        data_rows = rows
        header_row: List[OCRItem] = []
        desc_bounds = (0.0, float("inf"))
    else:
        header_row = rows[header_idx]
        data_rows = rows[header_idx + 1 :]
        max_x = max((item.x_max for row in data_rows for item in row), default=0.0)
        desc_bounds = compute_description_bounds(header_row, max_x)
    po_order: List[int] = []
    po_map: Dict[int, str] = {}
    for row in data_rows:
        left_candidates = [item for item in row if item.x_center <= desc_bounds[0]]
        po_line_value: Optional[int] = None
        for item in left_candidates:
            digits = regex.sub(r"[^0-9]", "", item.text)
            if digits and len(digits) <= 3:
                value = int(digits)
                if 0 < value < 1000:
                    po_line_value = value
                    break
        if po_line_value is None:
            continue
        desc_items = [
            item
            for item in row
            if desc_bounds[0] - 5 <= item.x_center <= desc_bounds[1] + 5
        ]
        if not desc_items:
            desc_items = [item for item in row if item.x_center > desc_bounds[0]]
        description = normalize_text(" ".join(item.text for item in desc_items))
        if po_line_value not in po_map:
            po_map[po_line_value] = description
            po_order.append(po_line_value)
    logger.log(f"Captured {len(po_order)} PO lines from OCR.")
    return po_order, po_map


def build_existing_description_order(ws) -> List[str]:
    values = []
    for row in range(4, 22):
        value = ws.cell(row=row, column=1).value
        values.append(normalize_text(value) if value else "")
    return values


def populate_po_descriptions(ws, po_order: List[int], po_map: Dict[int, str], logger: BuildLogger) -> Dict[int, Tuple[int, int]]:
    mapping: Dict[int, Tuple[int, int]] = {}
    existing = build_existing_description_order(ws)
    non_empty = [value for value in existing if value]
    if not non_empty and not DRY_RUN:
        logger.log("PO description rows are empty. Populating from OCR.")
        for idx, po_line in enumerate(po_order[:18]):
            description = po_map.get(po_line, "")
            top_row = 4 + idx
            ws.cell(row=top_row, column=1, value=description)
            ws.cell(row=25 + idx, column=1, value=description)
            mapping[po_line] = (top_row, 25 + idx)
        return mapping
    normalized_top: Dict[str, int] = {}
    for idx, desc in enumerate(existing):
        if desc:
            normalized_top[normalize_for_match(desc)] = 4 + idx
    normalized_matrix: Dict[str, int] = {}
    for offset in range(18):
        desc = ws.cell(row=25 + offset, column=1).value
        if desc:
            normalized_matrix[normalize_for_match(desc)] = 25 + offset
    for po_line, description in po_map.items():
        norm = normalize_for_match(description)
        top_row = normalized_top.get(norm)
        matrix_row = normalized_matrix.get(norm)
        if top_row is None and norm in normalized_matrix:
            top_row = 4 + (normalized_matrix[norm] - 25)
        if matrix_row is None and top_row is not None:
            matrix_row = 25 + (top_row - 4)
        if top_row is not None and matrix_row is not None:
            mapping[po_line] = (top_row, matrix_row)
        else:
            logger.log(
                f"Unable to align PO line {po_line} ('{po_map[po_line]}') with existing rows."
            )
    return mapping


def detect_po_number(all_text: str, logger: BuildLogger) -> Optional[str]:
    pattern = regex.compile(
        r"(\d+\.[A-Za-z]\.[A-Za-z]\.[A-Za-z]\.[0-9]+\.R|[0-9A-Za-z\.\-_/]*\.R)",
        regex.IGNORECASE,
    )
    matches = [normalize_text(match) for match in pattern.findall(all_text) if match]
    if matches:
        logger.log(f"Detected PO number candidate '{matches[0]}'.")
        return matches[0]
    logger.log("No PO number detected from OCR text.")
    return None


def group_header_tokens(header_row: List[OCRItem], expected_columns: int) -> List[Tuple[float, float]]:
    if not header_row:
        return []
    tokens = sorted(header_row, key=lambda item: item.x_center)
    groups: List[List[OCRItem]] = []
    current = [tokens[0]]
    for item in tokens[1:]:
        if item.x_center - current[-1].x_center <= 50:
            current.append(item)
        else:
            groups.append(current)
            current = [item]
    groups.append(current)
    while len(groups) > expected_columns:
        smallest_gap = None
        merge_idx = None
        for idx in range(len(groups) - 1):
            gap = groups[idx + 1][0].x_center - groups[idx][-1].x_center
            if smallest_gap is None or gap < smallest_gap:
                smallest_gap = gap
                merge_idx = idx
        if merge_idx is None:
            break
        groups[merge_idx].extend(groups[merge_idx + 1])
        del groups[merge_idx + 1]
    while len(groups) < expected_columns and groups:
        groups.append(groups[-1])
    boundaries: List[Tuple[float, float]] = []
    for group in groups:
        left = min(item.x_min for item in group)
        right = max(item.x_max for item in group)
        boundaries.append((left, right))
    adjusted: List[Tuple[float, float]] = []
    for idx, (left, right) in enumerate(boundaries):
        prev_right = boundaries[idx - 1][1] if idx > 0 else left
        next_left = boundaries[idx + 1][0] if idx + 1 < len(boundaries) else right
        adj_left = (prev_right + left) / 2 if idx > 0 else left - 20
        adj_right = (right + next_left) / 2 if idx + 1 < len(boundaries) else right + 200
        if adj_left >= adj_right:
            adj_left = left - 20
            adj_right = right + 20
        adjusted.append((adj_left, adj_right))
    return adjusted


def find_invoice_header_index(rows: List[List[OCRItem]]) -> Optional[int]:
    for idx, row in enumerate(rows):
        text = " ".join(item.text.casefold() for item in row)
        if "invoice" in text and "po" in text and "gross" in text:
            return idx
    return None


def assign_tokens_to_columns(
    row: List[OCRItem],
    boundaries: List[Tuple[float, float]],
    labels: List[str],
) -> Tuple[Dict[str, List[str]], List[float]]:
    values: Dict[str, List[str]] = {label: [] for label in labels}
    confidences: List[float] = []
    if not row:
        return values, confidences
    for item in sorted(row, key=lambda it: it.x_center):
        index = None
        for idx, (left, right) in enumerate(boundaries):
            if left <= item.x_center <= right:
                index = idx
                break
        if index is None:
            if item.x_center < boundaries[0][0]:
                index = 0
            else:
                index = len(boundaries) - 1
        values[labels[index]].append(item.text)
        confidences.append(item.confidence)
    return values, confidences


def parse_currency(value: str) -> Optional[float]:
    if not value:
        return None
    cleaned = value.replace("$", "").replace(",", "").replace(" ", "")
    cleaned = cleaned.replace("(", "-").replace(")", "")
    if cleaned in {"", "-", "--"}:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def parse_invoice_rows(
    rows: List[List[OCRItem]],
    logger: BuildLogger,
) -> Tuple[List[InvoiceRow], int]:
    header_idx = find_invoice_header_index(rows)
    if header_idx is None:
        logger.log("Invoice log header not detected; using first row as header.")
        header_idx = 0 if rows else None
    if header_idx is None:
        return [], 0
    header_row = rows[header_idx]
    data_rows = rows[header_idx + 1 :]
    column_labels = [
        "Status",
        "Invoice #",
        "PO Line",
        "Units",
        "Unit Cost",
        "Gross",
        "Description",
        "Tax",
    ]
    boundaries = group_header_tokens(header_row, len(column_labels))
    kept_rows: List[InvoiceRow] = []
    dropped = 0
    for row in data_rows:
        column_values, confs = assign_tokens_to_columns(row, boundaries, column_labels)
        avg_conf = statistics.mean(confs) if confs else 0.0
        values = {label: normalize_text(" ".join(parts)) for label, parts in column_values.items()}
        gross = parse_currency(values.get("Gross", ""))
        tax = parse_currency(values.get("Tax", ""))
        po_line = None
        po_line_text = values.get("PO Line", "")
        digits = regex.findall(r"\d+", po_line_text)
        if digits:
            try:
                po_line = int(digits[0])
            except ValueError:
                po_line = None
        invoice_number = values.get("Invoice #", "")
        if gross is None or avg_conf < 0.60 or not invoice_number:
            dropped += 1
            continue
        kept_rows.append(
            InvoiceRow(
                invoice_number=invoice_number,
                po_line=po_line,
                gross=gross,
                tax=tax,
                units=values.get("Units", ""),
                unit_cost=values.get("Unit Cost", ""),
                status=values.get("Status", ""),
                description=values.get("Description", ""),
                confidence=avg_conf,
            )
        )
    logger.log(f"Parsed {len(kept_rows)} invoice rows (dropped {dropped}).")
    return kept_rows, dropped


def save_invoice_csv(rows: List[InvoiceRow]):
    with CSV_PATH.open("w", newline="", encoding="utf-8") as fp:
        writer = csv.writer(fp)
        writer.writerow([
            "Invoice #",
            "PO Line",
            "Gross",
            "Tax",
            "Units",
            "Unit Cost",
            "Status",
            "Description",
        ])
        for row in rows:
            writer.writerow(
                [
                    row.invoice_number,
                    row.po_line if row.po_line is not None else "",
                    row.gross if row.gross is not None else "",
                    row.tax if row.tax is not None else "",
                    row.units,
                    row.unit_cost,
                    row.status,
                    row.description,
                ]
            )


def ensure_tax_rate(ws, logger: BuildLogger):
    cell_d1 = ws.cell(row=1, column=4)
    cell_e1 = ws.cell(row=1, column=5)
    if normalize_text(cell_d1.value).casefold() == "tax rate" and (cell_e1.value in (None, "")):
        if not DRY_RUN:
            cell_e1.value = TAX_RATE_DEFAULT
        cell_e1.number_format = "0.0000"
        logger.log("Set default tax rate in E1.")


def ensure_po_number(ws, po_number: Optional[str], logger: BuildLogger):
    if not po_number:
        return
    cell_b1 = ws.cell(row=1, column=2)
    if cell_b1.value in (None, "") and not DRY_RUN:
        cell_b1.value = po_number
    cell_b1.number_format = TEXT_FORMAT
    logger.log(f"PO number set to '{po_number}'.")


def build_po_mapping(
    ws,
    po_order: List[int],
    mapping_from_population: Dict[int, Tuple[int, int]],
) -> Dict[int, Tuple[int, int]]:
    mapping = mapping_from_population.copy()
    if not mapping:
        for idx, po_line in enumerate(po_order[:18]):
            mapping[po_line] = (4 + idx, 25 + idx)
    return mapping


def build_invoice_columns(ws, invoices: List[str], logger: BuildLogger) -> Dict[str, int]:
    if not invoices:
        return {}
    # existing headers
    existing_cols: Dict[str, int] = {}
    filled_cols = []
    max_col = ws.max_column
    for col in range(2, max_col + 1):
        value = ws.cell(row=24, column=col).value
        if value not in (None, ""):
            header_value = normalize_text(str(value))
            existing_cols[header_value] = col
            filled_cols.append(col)
            ws.cell(row=24, column=col).number_format = TEXT_FORMAT
    next_col = max(filled_cols) + 1 if filled_cols else 2
    assigned: Dict[str, int] = {}
    for invoice in invoices:
        normalized = normalize_text(invoice)
        if normalized in existing_cols:
            assigned[invoice] = existing_cols[normalized]
            continue
        col_idx = next_col
        if not DRY_RUN:
            cell = ws.cell(row=24, column=col_idx)
            cell.value = invoice
            cell.number_format = TEXT_FORMAT
        assigned[invoice] = col_idx
        next_col += 1
    if not DRY_RUN and ws.cell(row=24, column=1).value in (None, ""):
        ws.cell(row=24, column=1, value="INV")
    logger.log(f"Prepared {len(assigned)} invoice column headers.")
    return assigned


def matrix_description_map(ws) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for offset in range(18):
        desc = ws.cell(row=25 + offset, column=1).value
        if desc:
            mapping[normalize_for_match(desc)] = 25 + offset
    return mapping


def place_invoice_values(
    ws,
    invoice_rows: List[InvoiceRow],
    invoice_columns: Dict[str, int],
    po_mapping: Dict[int, Tuple[int, int]],
    logger: BuildLogger,
) -> List[str]:
    unmatched: List[str] = []
    matrix_map = matrix_description_map(ws)
    for row in invoice_rows:
        invoice_col = invoice_columns.get(row.invoice_number)
        if invoice_col is None:
            continue
        matrix_row = None
        if row.po_line is not None:
            matrix_info = po_mapping.get(row.po_line)
            if matrix_info:
                matrix_row = matrix_info[1]
        if matrix_row is None and row.description:
            norm = normalize_for_match(row.description)
            matrix_row = matrix_map.get(norm)
        if matrix_row is None:
            unmatched.append(
                f"Invoice {row.invoice_number} PO line {row.po_line} description '{row.description}'"
            )
            continue
        cell = ws.cell(row=matrix_row, column=invoice_col)
        if not DRY_RUN:
            existing_value = cell.value if isinstance(cell.value, (int, float)) else 0
            cell.value = (existing_value or 0) + (row.gross or 0)
            cell.number_format = CURRENCY_FORMAT
            cell.fill = MATRIX_FILL
    return unmatched


def update_formulas(
    ws,
    invoice_columns: Dict[str, int],
    logger: BuildLogger,
):
    if not invoice_columns:
        return
    last_col_index = max(invoice_columns.values())
    last_col_letter = get_column_letter(last_col_index)
    for idx in range(18):
        top_row = 4 + idx
        matrix_row = 25 + idx
        cell_e = ws.cell(row=top_row, column=5)
        if cell_e.value in (None, ""):
            if not DRY_RUN:
                cell_e.value = f"=SUM(B{matrix_row}:{last_col_letter}{matrix_row})"
            cell_e.number_format = CURRENCY_FORMAT
        else:
            logger.log(f"Skipped writing Total formula in E{top_row} (cell already populated).")
        cell_g = ws.cell(row=top_row, column=7)
        if cell_g.value in (None, ""):
            if not DRY_RUN:
                cell_g.value = f"=C{top_row}-E{top_row}"
        else:
            logger.log(f"Skipped writing Delta formula in G{top_row} (cell already populated).")
    for invoice, col_idx in invoice_columns.items():
        col_letter = get_column_letter(col_idx)
        if not DRY_RUN:
            ws.cell(row=44, column=col_idx).value = f"=SUM({col_letter}25:{col_letter}42)"
            ws.cell(row=45, column=col_idx).value = (
                f"=SUMIF(Invoice_Log!$A:$A,{col_letter}$24,Invoice_Log!$C:$C)"
            )
            ws.cell(row=46, column=col_idx).value = f"={col_letter}44-{col_letter}45"
        ws.cell(row=44, column=col_idx).number_format = CURRENCY_FORMAT
        ws.cell(row=45, column=col_idx).number_format = CURRENCY_FORMAT
        ws.cell(row=46, column=col_idx).number_format = CURRENCY_FORMAT
        formula = f"=ABS({col_letter}46)>0.01"
        ws.conditional_formatting.add(
            f"{col_letter}46",
            FormulaRule(formula=[formula], fill=VARIANCE_FILL),
        )
    logger.log("Updated summary formulas for totals and variances.")


def refresh_invoice_log_sheet(wb, invoice_rows: List[InvoiceRow]):
    if "Invoice_Log" in wb.sheetnames:
        sheet = wb["Invoice_Log"]
        wb.remove(sheet)
    sheet = wb.create_sheet("Invoice_Log")
    headers = [
        "Invoice #",
        "PO Line",
        "Gross",
        "Tax",
        "Units",
        "Unit Cost",
        "Status",
        "Description",
        "Confidence",
    ]
    for idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True)
    for row_idx, invoice_row in enumerate(invoice_rows, start=2):
        sheet.cell(row=row_idx, column=1, value=invoice_row.invoice_number)
        sheet.cell(row=row_idx, column=2, value=invoice_row.po_line)
        sheet.cell(row=row_idx, column=3, value=invoice_row.gross)
        sheet.cell(row=row_idx, column=4, value=invoice_row.tax)
        sheet.cell(row=row_idx, column=5, value=invoice_row.units)
        sheet.cell(row=row_idx, column=6, value=invoice_row.unit_cost)
        sheet.cell(row=row_idx, column=7, value=invoice_row.status)
        sheet.cell(row=row_idx, column=8, value=invoice_row.description)
        sheet.cell(row=row_idx, column=9, value=round(invoice_row.confidence, 4))
        sheet.cell(row=row_idx, column=3).number_format = CURRENCY_FORMAT
        sheet.cell(row=row_idx, column=4).number_format = CURRENCY_FORMAT
    sheet.freeze_panes = "A2"


def gather_all_text(items: List[OCRItem]) -> str:
    return " \n".join(item.text for item in items)


def write_log_summary(
    logger: BuildLogger,
    total_rows: int,
    dropped_rows: int,
    invoice_columns: Dict[str, int],
    unmatched: List[str],
    ctx: WorkbookContext,
):
    kept = total_rows - dropped_rows
    if total_rows == 0:
        kept = 0
    logger.log(f"Invoice rows kept: {kept}; dropped: {dropped_rows}.")
    logger.log(f"Invoice column count: {len(invoice_columns)}.")
    if invoice_columns:
        first_col = get_column_letter(min(invoice_columns.values()))
        last_col = get_column_letter(max(invoice_columns.values()))
        logger.log(f"Invoice headers span columns {first_col} to {last_col}.")
    if unmatched:
        logger.log("Unmatched invoice rows:")
        for entry in unmatched:
            logger.log(f"  - {entry}")
    if ctx.template_used:
        logger.log("Used existing template workbook.")
    if ctx.template_created:
        logger.log("Created new template workbook.")


def main():
    logger = BuildLogger()
    ctx = load_or_create_workbook(logger)
    ws = ctx.sheet
    reader = easyocr.Reader(["en"], gpu=False)

    po_results = read_image_text(PO_IMAGE_PATH, reader, logger)
    po_items = convert_results_to_items(po_results)
    po_rows = group_items_by_row(po_items)
    po_order, po_map = extract_po_lines(po_rows, logger)

    populated_mapping = populate_po_descriptions(ws, po_order, po_map, logger)
    po_mapping = build_po_mapping(ws, po_order, populated_mapping)

    invoice_results = read_image_text(INVOICE_IMAGE_PATH, reader, logger)
    invoice_items = convert_results_to_items(invoice_results)
    invoice_rows_grouped = group_items_by_row(invoice_items, tolerance=14.0)
    invoice_rows, dropped_rows = parse_invoice_rows(invoice_rows_grouped, logger)
    save_invoice_csv(invoice_rows)

    ensure_tax_rate(ws, logger)
    combined_text = gather_all_text(po_items + invoice_items)
    po_number = detect_po_number(combined_text, logger)
    ensure_po_number(ws, po_number, logger)

    invoices_in_order: List[str] = []
    seen_invoices = set()
    for row in invoice_rows:
        if row.invoice_number not in seen_invoices:
            invoices_in_order.append(row.invoice_number)
            seen_invoices.add(row.invoice_number)
    invoice_columns = build_invoice_columns(ws, invoices_in_order, logger)
    unmatched_rows = place_invoice_values(ws, invoice_rows, invoice_columns, logger)
    update_formulas(ws, invoice_columns, logger)
    refresh_invoice_log_sheet(ctx.workbook, invoice_rows)

    ctx.workbook.save(OUTPUT_PATH)
    logger.log(f"Workbook saved to {OUTPUT_PATH.name}.")
    write_log_summary(logger, len(invoice_rows) + dropped_rows, dropped_rows, invoice_columns, unmatched_rows, ctx)
    logger.write(LOG_PATH)

    print("PO number:", po_number or "(not found)")
    print("Invoice columns:", len(invoice_columns))
    print("Output workbook:", OUTPUT_PATH)


if __name__ == "__main__":
    main()
